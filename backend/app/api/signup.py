"""API auto-signup."""

from __future__ import annotations

import io
import json
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.db import get_session
from app.deps import get_current_user
from app.models import User
from app.models.signup_job import SignupJob
from app.services import signup_runner

router = APIRouter(prefix="/api/signup", tags=["signup"])


class SignupJobIn(BaseModel):
    program_ids: List[int]
    profile_ids: List[str]
    email_ids: List[str] = []
    proxy_ids: List[str] = []
    instruction_names: List[str] = []
    instruction_name: Optional[str] = ""  # legacy single
    extra_prompt: Optional[str] = ""
    headless: bool = False
    gemini_key_index: Optional[int] = 0  # legacy; still accepted for old clients
    llm_provider: Optional[str] = ""   # "gemini"|"openai"|"deepseek"|"" = auto
    llm_key_index: Optional[int] = 0   # 1-based; 0/None = auto-rotate
    run_mode: Optional[str] = "llm"    # "llm" | "script" | "script_llm"
    playbook_id: Optional[int] = None  # playbook to use when run_mode in ("script","script_llm")
    sms_profile_id: Optional[str] = ""  # áp preset SMS cho cả job (override per-program)
    batch_id: Optional[str] = ""  # gom nhiều job vào 1 batch
    # Script Signup v2 fields
    tier3_behavior: Optional[str] = "ask"   # "ask" | "auto" | "never"
    script_overrides: Optional[Dict[str, Any]] = None  # {program_id: playbook_id}


class SignupJobOut(BaseModel):
    id: int
    user_id: Optional[int]
    program_ids: List[int]
    profile_ids: List[str]
    email_ids: List[str] = []
    proxy_ids: List[str] = []
    instruction_names: List[str] = []
    instruction_name: Optional[str]
    extra_prompt: Optional[str]
    headless: bool
    status: str
    total: int
    succeeded: int
    failed: int
    results: list
    error: Optional[str]
    started_at: Optional[str]
    finished_at: Optional[str]
    created_at: Optional[str]
    sms_profile_id: Optional[str] = None
    gemini_key_index: Optional[int] = None
    llm_provider: Optional[str] = None
    llm_key_index: Optional[int] = None
    run_mode: Optional[str] = "llm"
    playbook_id: Optional[int] = None
    batch_id: Optional[str] = None
    tier3_behavior: Optional[str] = "ask"
    tier3_status: Optional[str] = None
    script_overrides: Optional[Dict[str, Any]] = None


def _to_out(j: SignupJob) -> SignupJobOut:
    return SignupJobOut(
        id=j.id,
        user_id=j.user_id,
        program_ids=json.loads(j.program_ids_json or "[]"),
        profile_ids=json.loads(j.profile_ids_json or "[]"),
        email_ids=json.loads(j.email_ids_json or "[]"),
        proxy_ids=json.loads(j.proxy_ids_json or "[]"),
        instruction_names=json.loads(j.instruction_names_json or "[]"),
        instruction_name=j.instruction_name,
        extra_prompt=j.extra_prompt,
        headless=j.headless,
        status=j.status,
        total=j.total,
        succeeded=j.succeeded,
        failed=j.failed,
        results=json.loads(j.results_json or "[]"),
        error=j.error,
        started_at=j.started_at.isoformat() + "Z" if j.started_at else None,
        finished_at=j.finished_at.isoformat() + "Z" if j.finished_at else None,
        created_at=j.created_at.isoformat() + "Z" if j.created_at else None,
        sms_profile_id=j.sms_profile_id,
        gemini_key_index=j.gemini_key_index,
        llm_provider=j.llm_provider,
        llm_key_index=j.llm_key_index,
        run_mode=j.run_mode or "llm",
        playbook_id=j.playbook_id,
        batch_id=j.batch_id,
        tier3_behavior=j.tier3_behavior or "ask",
        tier3_status=j.tier3_status,
        script_overrides=json.loads(j.script_overrides_json or "{}") or None,
    )


@router.post("/jobs", response_model=SignupJobOut)
async def create_signup_job(
    payload: SignupJobIn,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_session),
):
    if not payload.program_ids:
        raise HTTPException(400, "Cần chọn ít nhất 1 program")
    if not payload.profile_ids:
        raise HTTPException(400, "Cần chọn ít nhất 1 profile")
    job_id = await signup_runner.enqueue(
        user_id=user.id,
        program_ids=payload.program_ids,
        profile_ids=payload.profile_ids,
        email_ids=payload.email_ids,
        proxy_ids=payload.proxy_ids,
        instruction_names=payload.instruction_names,
        instruction_name=payload.instruction_name or "",
        extra_prompt=payload.extra_prompt or "",
        headless=payload.headless,
        sms_profile_id=payload.sms_profile_id or "",
        gemini_key_index=payload.gemini_key_index or 0,
        llm_provider=payload.llm_provider or "",
        llm_key_index=payload.llm_key_index or 0,
        run_mode=payload.run_mode or "llm",
        playbook_id=payload.playbook_id,
        batch_id=payload.batch_id or "",
        tier3_behavior=payload.tier3_behavior or "ask",
        script_overrides=payload.script_overrides or {},
    )
    row = (
        await session.execute(select(SignupJob).where(SignupJob.id == job_id))
    ).scalar_one()
    return _to_out(row)


@router.get("/jobs", response_model=List[SignupJobOut])
async def list_signup_jobs(
    limit: int = 50,
    run_mode: Optional[str] = None,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_session),
):
    q = select(SignupJob).where(SignupJob.user_id == user.id)
    if run_mode == "llm":
        q = q.where(or_(SignupJob.run_mode.is_(None), SignupJob.run_mode == "", SignupJob.run_mode == "llm"))
    elif run_mode == "script":
        q = q.where(SignupJob.run_mode.in_(["script", "script_llm"]))
    rows = (await session.execute(q.order_by(SignupJob.id.desc()).limit(limit))).scalars().all()
    return [_to_out(r) for r in rows]


@router.get("/jobs/{job_id}", response_model=SignupJobOut)
async def get_signup_job(
    job_id: int,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_session),
):
    row = (
        await session.execute(select(SignupJob).where(SignupJob.id == job_id))
    ).scalar_one_or_none()
    if not row:
        raise HTTPException(404, "Job không tồn tại")
    if row.user_id and row.user_id != user.id:
        raise HTTPException(403, "Không có quyền xem job này")
    return _to_out(row)


@router.patch("/jobs/{job_id}/cancel", response_model=SignupJobOut)
async def cancel_signup_job(
    job_id: int,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_session),
):
    """Cancel a pending or running signup job."""
    from datetime import datetime
    row = (
        await session.execute(select(SignupJob).where(SignupJob.id == job_id))
    ).scalar_one_or_none()
    if not row:
        raise HTTPException(404, "Job không tồn tại")
    if row.user_id and row.user_id != user.id:
        raise HTTPException(403, "Không có quyền")
    if row.status not in ("pending", "running"):
        raise HTTPException(400, f"Không thể cancel job ở trạng thái {row.status}")
    row.status = "failed"
    row.error = "Cancelled by user"
    row.finished_at = datetime.utcnow()
    await session.commit()
    await session.refresh(row)
    return _to_out(row)


@router.post("/jobs/{job_id}/approve-tier3", response_model=SignupJobOut)
async def approve_tier3(
    job_id: int,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_session),
):
    """User duyệt chạy Tier 3 cho các program đang chờ approval trong job này.
    Tạo job mới (run_mode=llm) cho những program có status='tier3_waiting'.
    """
    from datetime import datetime
    row = (
        await session.execute(select(SignupJob).where(SignupJob.id == job_id))
    ).scalar_one_or_none()
    if not row:
        raise HTTPException(404, "Job không tồn tại")
    if row.user_id and row.user_id != user.id:
        raise HTTPException(403, "Không có quyền")
    if row.tier3_status != "waiting_approval":
        raise HTTPException(400, "Job này không có Tier 3 đang chờ duyệt")

    # Lấy danh sách program đang tier3_waiting
    results = json.loads(row.results_json or "[]")
    waiting_ids = [r["program_id"] for r in results if r.get("status") == "tier3_waiting"]
    if not waiting_ids:
        raise HTTPException(400, "Không tìm thấy program nào đang chờ Tier 3")

    # Tạo job mới với run_mode=llm (Tier 3 full AI)
    new_job_id = await signup_runner.enqueue(
        user_id=user.id,
        program_ids=waiting_ids,
        profile_ids=json.loads(row.profile_ids_json or "[]"),
        email_ids=json.loads(row.email_ids_json or "[]"),
        proxy_ids=json.loads(row.proxy_ids_json or "[]"),
        instruction_names=json.loads(row.instruction_names_json or "[]"),
        instruction_name=row.instruction_name or "",
        extra_prompt=row.extra_prompt or "",
        headless=bool(row.headless),
        sms_profile_id=row.sms_profile_id or "",
        llm_provider=row.llm_provider or "",
        llm_key_index=int(row.llm_key_index or 0),
        run_mode="llm",
        tier3_behavior="auto",
        script_overrides={},
    )

    # Cập nhật job gốc
    row.tier3_status = "approved"
    await session.commit()
    await session.refresh(row)
    return _to_out(row)


@router.post("/jobs/{job_id}/skip-tier3", response_model=SignupJobOut)
async def skip_tier3(
    job_id: int,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_session),
):
    """User bỏ qua Tier 3 — đánh dấu các program tier3_waiting là skipped."""
    from datetime import datetime
    row = (
        await session.execute(select(SignupJob).where(SignupJob.id == job_id))
    ).scalar_one_or_none()
    if not row:
        raise HTTPException(404, "Job không tồn tại")
    if row.user_id and row.user_id != user.id:
        raise HTTPException(403, "Không có quyền")
    if row.tier3_status != "waiting_approval":
        raise HTTPException(400, "Job này không có Tier 3 đang chờ duyệt")

    results = json.loads(row.results_json or "[]")
    for r in results:
        if r.get("status") == "tier3_waiting":
            r["status"] = "skipped"
            r["message"] = "Tier 3 bị bỏ qua bởi user"

    row.tier3_status = "skipped"
    row.results_json = json.dumps(results, ensure_ascii=False)
    failed_count = sum(1 for r in results if r.get("status") not in ("success", "pending_verify"))
    succeeded_count = sum(1 for r in results if r.get("status") in ("success", "pending_verify"))
    row.failed = failed_count
    row.succeeded = succeeded_count
    await session.commit()
    await session.refresh(row)
    return _to_out(row)


@router.get("/screenshots/{filename:path}")
async def get_signup_screenshot(
    filename: str,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_session),
):
    """Serve PNG screenshot từ data/signup_screenshots/."""
    if ".." in filename or filename.startswith("/"):
        raise HTTPException(400, "Invalid filename")
    # Filename có thể bao gồm prefix "signup_screenshots/" do backend lưu relative path
    name = filename.split("/")[-1]
    if not name.endswith(".png"):
        raise HTTPException(400, "Chỉ phục vụ file .png")
    path = settings.data_path("signup_screenshots") / name
    if not path.exists():
        raise HTTPException(404, "Screenshot không tồn tại")
    # Verify ownership: parse job_id từ filename (format: job{id}_prog...) → check user
    import re as _re
    m = _re.match(r"job(\d+)_", name)
    if m:
        job_id = int(m.group(1))
        row = (await session.execute(
            select(SignupJob).where(SignupJob.id == job_id)
        )).scalar_one_or_none()
        if row and row.user_id and row.user_id != user.id:
            raise HTTPException(403, "Không có quyền xem screenshot này")
    return FileResponse(path, media_type="image/png")


# ─────────────────────────────────────────────────────────
# Helpers shared by /history and /history/export
# ─────────────────────────────────────────────────────────

async def _build_history_items(
    user: User,
    session: AsyncSession,
    status_filter: str = "",
    search: str = "",
    date_from: str = "",
    date_to: str = "",
    job_ids: str = "",
) -> List[Dict[str, Any]]:
    """Flatten tất cả signup attempts, enrich với program + profile data."""
    from app.models.affiliate_program import AffiliateProgram
    from app.services.storage import profile_store as _ps

    # 1. Load all jobs
    rows = (
        await session.execute(
            select(SignupJob)
            .where(SignupJob.user_id == user.id)
            .order_by(SignupJob.created_at.desc())
        )
    ).scalars().all()

    # 2. Flatten attempts
    allowed_job_ids: set[int] | None = None
    if job_ids.strip():
        try:
            allowed_job_ids = {int(x) for x in job_ids.split(",") if x.strip()}
        except ValueError:
            pass
    raw: List[Dict[str, Any]] = []
    for job in rows:
        if allowed_job_ids is not None and job.id not in allowed_job_ids:
            continue
        attempts = json.loads(job.results_json or "[]")
        for attempt in attempts:
            if isinstance(attempt, dict):
                raw.append({"job": job, "attempt": attempt})

    # 3. Bulk load programs
    prog_ids = list({a["attempt"].get("program_id") for a in raw if a["attempt"].get("program_id")})
    programs: Dict[int, Any] = {}
    if prog_ids:
        for p in (await session.execute(
            select(AffiliateProgram).where(AffiliateProgram.id.in_(prog_ids))
        )).scalars().all():
            programs[p.id] = p

    # 4. Bulk load profiles
    profile_ids = list({a["attempt"].get("profile_id") for a in raw if a["attempt"].get("profile_id")})
    profiles: Dict[str, Dict] = {}
    for pid in profile_ids:
        pd = _ps.get_profile(user.id, pid)
        if pd:
            profiles[pid] = pd

    # 5. Build items + apply filters
    items: List[Dict[str, Any]] = []
    for row in raw:
        job = row["job"]
        attempt = row["attempt"]
        prog_id = attempt.get("program_id")
        prog = programs.get(prog_id) if prog_id else None
        profile_id = attempt.get("profile_id") or ""
        profile = profiles.get(profile_id) or {}

        attempt_status = (attempt.get("status") or "unknown").lower()
        prog_name = prog.name if prog else f"Program #{prog_id}"
        prog_source = (prog.source if prog else "") or ""
        started_at = attempt.get("started_at") or (
            job.started_at.isoformat() + "Z" if job.started_at else None
        )

        # Filters
        if status_filter and attempt_status != status_filter.lower():
            continue
        if search and search.lower() not in prog_name.lower():
            continue
        if date_from and started_at and started_at[:10] < date_from:
            continue
        if date_to and started_at and started_at[:10] > date_to:
            continue

        items.append({
            "job_id": job.id,
            "job_created_at": job.created_at.isoformat() + "Z" if job.created_at else None,
            "program_id": prog_id,
            "program_name": prog_name,
            "program_url": prog.url if prog else None,
            "program_signup_url": prog.signup_url if prog else attempt.get("final_url"),
            "program_commission": prog.commission if prog else None,
            "program_source": prog_source,
            "program_logo_url": getattr(prog, "logo_url", None),
            "profile_id": profile_id,
            "profile_name": profile.get("full_name") or profile_id,
            "profile_email": profile.get("email") or "",
            "profile_website": profile.get("website") or "",
            "profile_niche": profile.get("niche") or [],
            # Actual data used at runtime (saved in signup_runner.py entry)
            "email_used": attempt.get("email_used") or "",
            "proxy_label": attempt.get("proxy_label") or "",
            "proxy_host": attempt.get("proxy_host") or "",
            "proxy_url_used": attempt.get("proxy_url_used") or "",
            "instruction_names_used": attempt.get("instruction_names_used") or [],
            "profile_snapshot": attempt.get("profile_snapshot") or {},
            "status": attempt_status,
            "message": attempt.get("message") or "",
            "screenshot": attempt.get("screenshot") or "",
            "final_url": attempt.get("final_url") or "",
            "duration_sec": attempt.get("duration_sec"),
            "steps": attempt.get("steps"),
            "started_at": started_at,
            "finished_at": attempt.get("finished_at"),
        })

    return items


def _compute_stats(items: List[Dict[str, Any]]) -> Dict[str, Any]:
    status_counts = Counter(i["status"] for i in items)
    source_counts = Counter(i["program_source"] for i in items if i["program_source"])
    day_data: Dict[str, Dict] = {}
    for i in items:
        if i.get("started_at"):
            day = i["started_at"][:10]
            bucket = day_data.setdefault(day, {"date": day, "success": 0, "failed": 0, "total": 0})
            bucket["total"] += 1
            if i["status"] == "success":
                bucket["success"] += 1
            elif i["status"] in ("failed", "captcha", "error"):
                bucket["failed"] += 1
    return {
        "total": len(items),
        "success": status_counts.get("success", 0),
        "failed": status_counts.get("failed", 0) + status_counts.get("error", 0),
        "captcha": status_counts.get("captcha", 0),
        "pending_verify": status_counts.get("pending_verify", 0),
        "by_status": dict(status_counts),
        "by_source": dict(source_counts),
        "by_day": sorted(day_data.values(), key=lambda d: d["date"]),
    }


# ─────────────────────────────────────────────────────────
# GET /api/signup/history  — paginated history + stats
# ─────────────────────────────────────────────────────────

@router.get("/history")
async def get_signup_history(
    status: str = "",
    search: str = "",
    date_from: str = "",
    date_to: str = "",
    page: int = 1,
    limit: int = 50,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_session),
):
    """Lịch sử đăng ký chi tiết — tất cả attempts, kèm program + profile info và thống kê."""
    page = max(1, page)
    limit = max(1, min(200, limit))

    all_items = await _build_history_items(user, session, status, search, date_from, date_to)
    stats = _compute_stats(all_items)
    total_count = len(all_items)
    items_page = all_items[(page - 1) * limit: page * limit]

    return {
        "stats": stats,
        "items": items_page,
        "total_count": total_count,
        "page": page,
        "limit": limit,
    }


# ─────────────────────────────────────────────────────────
# GET /api/signup/history/export  — Excel download
# ─────────────────────────────────────────────────────────

@router.get("/history/export")
async def export_signup_history(
    status: str = "",
    search: str = "",
    date_from: str = "",
    date_to: str = "",
    job_ids: str = "",
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_session),
):
    """Xuất lịch sử đăng ký ra file Excel (.xlsx), có ảnh thumbnail."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter
    from datetime import datetime as _dt
    try:
        from PIL import Image as PILImage
        _has_pil = True
    except ImportError:
        _has_pil = False

    THUMB_W, THUMB_H = 200, 112   # 16:9 thumbnail (px)
    SCREENSHOT_COL = 16           # 1-based column index

    # ── Color palette ────────────────────────────────────────
    C_HEADER_BG   = "1E3A5F"
    C_HEADER_FG   = "FFFFFF"
    C_ALT_ROW     = "F7F9FC"
    C_BORDER      = "C8D0DC"
    C_URL         = "1A56DB"
    STATUS_CFG = {
        "success":        ("D1FAE5", "065F46", "✔ Thành công"),
        "failed":         ("FEE2E2", "991B1B", "✘ Thất bại"),
        "error":          ("FEE2E2", "991B1B", "⚠ Lỗi"),
        "captcha":        ("FEF3C7", "92400E", "🔒 Captcha"),
        "running":        ("DBEAFE", "1E40AF", "⏳ Đang chạy"),
        "pending_verify": ("EDE9FE", "5B21B6", "⏸ Chờ duyệt"),
    }

    def thin_border():
        s = Side(style="thin", color=C_BORDER)
        return Border(left=s, right=s, top=s, bottom=s)

    def make_url_font():
        return Font(name="Calibri", size=9, color=C_URL, underline="single")

    items = await _build_history_items(user, session, status, search, date_from, date_to, job_ids)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lịch sử đăng ký"

    # ── Headers ──────────────────────────────────────────────
    headers = [
        "STT", "Ngày gửi", "Tên chương trình", "Nguồn", "URL đăng ký",
        "Hoa hồng", "Tên profile", "Email dùng", "Website", "Niche",
        "Trạng thái", "Kết quả", "URL kết thúc", "T.gian (s)", "Bước",
        "Screenshot", "Job",
    ]
    # col widths (chars)
    col_widths = [5, 17, 30, 12, 38, 13, 20, 28, 25, 20, 15, 52, 38, 11, 7, 27, 6]
    # columns to center-align
    CENTER_COLS  = {1, 4, 6, 11, 14, 15, 17}
    # columns with URLs (show as hyperlink)
    URL_COLS     = {5, 13}

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font        = Font(name="Calibri", bold=True, size=10, color=C_HEADER_FG)
        cell.fill        = PatternFill(fill_type="solid", fgColor=C_HEADER_BG)
        cell.alignment   = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border      = thin_border()
    ws.row_dimensions[1].height = 24

    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ── Data rows ────────────────────────────────────────────
    for i, item in enumerate(items, 2):
        is_odd        = (i % 2 == 0)
        status_key    = (item["status"] or "").lower()
        st_cfg        = STATUS_CFG.get(status_key)
        row_bg        = st_cfg[0] if st_cfg else ("FFFFFF" if is_odd else C_ALT_ROW)

        date_str = ""
        if item.get("started_at"):
            date_str = item["started_at"][:19].replace("T", " ")

        screenshot_raw  = item["screenshot"] or ""
        screenshot_name = screenshot_raw.split("/")[-1] if screenshot_raw else ""
        email_used      = item.get("email_used") or item.get("profile_email") or ""

        row_data = [
            i - 1,
            date_str,
            item["program_name"],
            item["program_source"],
            item["program_signup_url"] or "",
            item["program_commission"] or "",
            item["profile_name"],
            email_used,
            item["profile_website"],
            ", ".join(item["profile_niche"] or []),
            st_cfg[2] if st_cfg else item["status"],          # human-readable status
            item["message"],
            item["final_url"] or "",
            round(item["duration_sec"], 1) if item.get("duration_sec") else "",
            item.get("steps") or "",
            screenshot_name,
            item["job_id"],
        ]

        for col, val in enumerate(row_data, 1):
            cell        = ws.cell(row=i, column=col, value=val)
            cell.fill   = PatternFill(fill_type="solid", fgColor=row_bg)
            cell.border = thin_border()
            h_align     = "center" if col in CENTER_COLS else "left"

            if col == 11 and st_cfg:
                # Status cell: colored bold text
                cell.font      = Font(name="Calibri", size=9, bold=True, color=st_cfg[1])
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col == 12:
                # Result/message: wrap text, small font
                cell.font      = Font(name="Calibri", size=8, color="374151")
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif col in URL_COLS and val:
                # Hyperlink
                cell.hyperlink = str(val)
                cell.font      = make_url_font()
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.font      = Font(name="Calibri", size=9)
                cell.alignment = Alignment(horizontal=h_align, vertical="center", wrap_text=False)

        # Embed screenshot thumbnail
        img_embedded = False
        if _has_pil and screenshot_name and screenshot_name.endswith(".png"):
            img_path = settings.data_path("signup_screenshots") / screenshot_name
            if img_path.exists():
                try:
                    with PILImage.open(img_path) as pil_img:
                        pil_img = pil_img.convert("RGB")
                        pil_img.thumbnail((THUMB_W, THUMB_H), PILImage.LANCZOS)
                        actual_w, actual_h = pil_img.size
                        thumb_buf = io.BytesIO()
                        pil_img.save(thumb_buf, format="PNG", optimize=True)
                        thumb_buf.seek(0)
                    xl_img         = XLImage(thumb_buf)
                    xl_img.width   = actual_w
                    xl_img.height  = actual_h
                    col_letter     = get_column_letter(SCREENSHOT_COL)
                    ws.cell(row=i, column=SCREENSHOT_COL, value="")
                    ws.add_image(xl_img, f"{col_letter}{i}")
                    ws.row_dimensions[i].height = max(actual_h * 0.75, 16)
                    img_embedded = True
                except Exception:
                    pass

        if not img_embedded:
            ws.row_dimensions[i].height = 18

    ws.freeze_panes = "C2"   # freeze STT+date, scroll program onwards

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"signup_history_{_dt.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
